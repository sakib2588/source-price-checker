"""The review workbook built from the store (owner layout: product, Source 1, Source 2, Reason)."""
import os
import tempfile
import unittest

import openpyxl

import sc_review_book
from test_sc_reuse import ReuseCase


class ReviewBook(ReuseCase):
    def test_the_workbook_has_the_owner_columns_and_one_row_per_product(self):
        self.check_everything()
        out = os.path.join(self.tmp.name, "review.xlsx")
        got = sc_review_book.build(self.store.path, out)
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Summary", "Products", "Details", "Sources", "About"])
        head = [c.value for c in wb["Products"][1]]
        self.assertEqual(head, ["Maskura product", "Source 1", "Source 2", "Reason"])
        self.assertEqual(wb["Products"].max_row - 1, got["products"])
